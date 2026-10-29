"""Genera ENTREGA-1-Iteracion-1.docx a partir de ENTREGA-1-Iteracion-1.md.

Uso:  <python> generar-docx.py <ruta-del-md> <ruta-del-docx>

Interpreta el subconjunto de Markdown que usa el documento: títulos, párrafos (unidos
aunque en el .md estén partidos en varias líneas), tablas, viñetas, listas numeradas,
imágenes, citas y bloques de código. La negrita (**texto**) y el código (`texto`) se
resuelven después de unir el párrafo, así que pueden cruzar de línea en el origen.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+`)")
BULLET = re.compile(r"^[-*]\s+(.*)$")
NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")
IMAGE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
SEPARADOR_TABLA = re.compile(r"^\|[\s:\-|]+\|$")


def agregar_runs(parrafo, texto, base_negrita=False, tamano=None):
    """Agrega texto a un párrafo interpretando **negrita** y `codigo`."""
    for parte in INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            run = parrafo.add_run(parte[1:-1])
            run.italic = True
        elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
            run = parrafo.add_run(parte[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            run = parrafo.add_run(parte)
        if base_negrita:
            run.bold = True
        if tamano is not None and run.font.size is None:
            run.font.size = Pt(tamano)


def numero_de_pagina(parrafo):
    run = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = "PAGE"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run._r.append(inicio)
    run._r.append(instruccion)
    run._r.append(fin)


def configurar(documento):
    normal = documento.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    for nivel in range(1, 4):
        estilo = documento.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    documento.styles["Heading 1"].font.size = Pt(16)
    documento.styles["Heading 2"].font.size = Pt(13)
    documento.styles["Heading 3"].font.size = Pt(11.5)

    pie = documento.sections[0].footer.paragraphs[0]
    pie.text = "SolarUY · Documentación del sistema · Equipo 6 · UTEC — página "
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    numero_de_pagina(pie)
    for run in pie.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)


class Generador:
    def __init__(self, documento, base):
        self.documento = documento
        self.base = base
        self.estado = None          # None | 'texto' | 'bullet' | 'numbered'
        self.buffer = []
        self.tabla = []
        self.en_codigo = False
        self.primer_h2 = True

    # ---------- salidas ----------
    def _texto_union(self):
        return " ".join(linea.strip() for linea in self.buffer if linea.strip())

    def emitir(self):
        if not self.buffer:
            self.estado = None
            return
        texto = self._texto_union()
        if self.estado == "bullet":
            parrafo = self.documento.add_paragraph(style="List Bullet")
            agregar_runs(parrafo, texto)
        elif self.estado == "numbered":
            parrafo = self.documento.add_paragraph()
            parrafo.paragraph_format.left_indent = Cm(0.6)
            parrafo.paragraph_format.first_line_indent = Cm(-0.6)
            agregar_runs(parrafo, texto)
        else:
            parrafo = self.documento.add_paragraph()
            agregar_runs(parrafo, texto)
        self.buffer = []
        self.estado = None

    def emitir_tabla(self):
        if not self.tabla:
            return
        filas = self.tabla
        self.tabla = []
        if all(not celda.strip() for celda in filas[0]):
            filas = filas[1:]
        if not filas:
            return
        ancho = max(len(f) for f in filas)
        tabla = self.documento.add_table(rows=0, cols=ancho)
        tabla.style = "Table Grid"
        for indice, fila in enumerate(filas):
            fila_docx = tabla.add_row()
            celdas = fila_docx.cells
            # Que una fila no se corte entre dos páginas.
            propiedades = OxmlElement("w:trPr")
            propiedades.append(OxmlElement("w:cantSplit"))
            fila_docx._tr.append(propiedades)
            for posicion in range(ancho):
                texto = fila[posicion] if posicion < len(fila) else ""
                parrafo = celdas[posicion].paragraphs[0]
                parrafo.paragraph_format.space_after = Pt(2)
                agregar_runs(parrafo, texto, base_negrita=(indice == 0), tamano=9.5)
        self.documento.add_paragraph()

    def cerrar_bloques(self):
        self.emitir()
        self.emitir_tabla()

    # ---------- una línea ----------
    def procesar(self, linea):
        limpia = linea.rstrip()
        despojada = limpia.strip()

        if limpia.startswith("```"):
            self.cerrar_bloques()
            self.en_codigo = not self.en_codigo
            return
        if self.en_codigo:
            parrafo = self.documento.add_paragraph(style="No Spacing")
            run = parrafo.add_run(limpia)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            return

        if limpia.startswith("|"):
            if not SEPARADOR_TABLA.match(limpia):
                self.tabla.append([c.strip() for c in limpia.strip("|").split("|")])
            return

        if despojada == "" or despojada == "---":
            self.cerrar_bloques()
            return

        if limpia.startswith("### "):
            self.cerrar_bloques()
            self.documento.add_heading(limpia[4:].strip(), level=3)
            return
        if limpia.startswith("## "):
            self.cerrar_bloques()
            if self.primer_h2:
                self.documento.add_page_break()
                self.primer_h2 = False
            self.documento.add_heading(limpia[3:].strip(), level=1)
            return
        if limpia.startswith("# "):
            self.cerrar_bloques()
            self.documento.add_heading(limpia[2:].strip(), level=0)
            return

        imagen = IMAGE.match(despojada)
        if imagen:
            self.cerrar_bloques()
            ruta = self.base / imagen.group(2)
            if ruta.exists():
                parrafo = self.documento.add_paragraph()
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Las figuras verticales se ajustan por alto para que entren en una página;
                # las horizontales, por ancho.
                from docx.image.image import Image as ImagenDocx

                medidas = ImagenDocx.from_file(str(ruta))
                if medidas.height > medidas.width:
                    parrafo.add_run().add_picture(str(ruta), height=Cm(18))
                else:
                    parrafo.add_run().add_picture(str(ruta), width=Cm(14.5))
            else:
                self.documento.add_paragraph(f"[falta la imagen {imagen.group(2)}]")
            return

        if despojada.startswith("> "):
            self.cerrar_bloques()
            parrafo = self.documento.add_paragraph(style="Intense Quote")
            agregar_runs(parrafo, despojada[2:].strip())
            return

        bullet = BULLET.match(despojada)
        if bullet:
            self.emitir()
            self.estado = "bullet"
            self.buffer = [bullet.group(1)]
            return

        numerada = NUMBERED.match(despojada)
        if numerada:
            self.emitir()
            self.estado = "numbered"
            self.buffer = [f"{numerada.group(1)}. {numerada.group(2)}"]
            return

        # Línea de texto: continúa el párrafo (o la viñeta) si corresponde.
        if self.estado in ("bullet", "numbered") and limpia[:1].isspace():
            self.buffer.append(limpia)
            return

        if self.estado == "texto":
            self.buffer.append(limpia)
        else:
            self.emitir()
            self.estado = "texto"
            self.buffer = [limpia]


def main():
    origen = Path(sys.argv[1])
    destino = Path(sys.argv[2])

    documento = Document()
    configurar(documento)
    generador = Generador(documento, origen.parent)

    for linea in origen.read_text(encoding="utf-8").splitlines():
        generador.procesar(linea)
    generador.cerrar_bloques()

    documento.save(destino)
    print(f"generado: {destino}")


if __name__ == "__main__":
    main()
